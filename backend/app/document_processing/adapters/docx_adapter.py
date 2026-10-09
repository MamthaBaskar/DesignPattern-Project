"""
DOCX Adapter implementing DocumentAdapterInterface.
Extracts paragraphs and tables from Microsoft Word (.docx) files.
"""

import io
import logging
from typing import Any, Dict, Tuple
import docx

from ..interfaces import DocumentAdapterInterface
from ..validators import DocumentValidationError

logger = logging.getLogger(__name__)


class DOCXAdapter(DocumentAdapterInterface):
    """
    Adapter that adapts python-docx to DocumentAdapterInterface.
    """

    def extract(self, file_bytes: bytes, filename: str) -> Tuple[str, Dict[str, Any]]:
        metadata: Dict[str, Any] = {
            "format": "docx",
            "paragraph_count": 0,
            "table_count": 0,
            "extraction_method": "docx_parser",
        }

        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise DocumentValidationError(
                f"Failed to read DOCX file '{filename}'. The file may be corrupted or not a valid Word document: {str(e)}",
                "CORRUPT_DOCX",
            )

        text_blocks = []

        # Extract regular paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                text_blocks.append(text)

        metadata["paragraph_count"] = len(text_blocks)

        # Extract text from tables
        table_count = 0
        for table in doc.tables:
            table_count += 1
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_blocks.append(" | ".join(row_cells))

        metadata["table_count"] = table_count

        extracted_text = "\n\n".join(text_blocks).strip()

        if not extracted_text:
            raise DocumentValidationError(
                f"DOCX document '{filename}' is empty or contains no readable text content.",
                "EMPTY_DOCX",
            )

        return extracted_text, metadata

